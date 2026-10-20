import io
import os
import re
import logging
import zipfile
import xml.etree.ElementTree as ET
from html.parser import HTMLParser
from typing import Tuple, Optional

logger = logging.getLogger(__name__)


# Learning Note:
# This module (Module 5) converts uploaded files (.pdf, .docx, .pptx, .xlsx, .html, .md, .txt, .csv, .json)
# into plain text/Markdown so they can be chunked and embedded across Legal, Inventory, KPI, and Client Presentation domains.
# It tries optional rich parsers first and includes zero-dependency Python standard library (zipfile + XML) fallbacks.

def try_docling_parse(filename: str, content_bytes: bytes) -> Optional[str]:
    """
    Attempt to convert document using IBM Docling (DocumentConverter) if installed.
    Exports structured Markdown layout.
    """
    try:
        from docling.document_converter import DocumentConverter
        import tempfile

        ext = os.path.splitext(filename)[1].lower()
        with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
            tmp.write(content_bytes)
            tmp_path = tmp.name

        try:
            converter = DocumentConverter()
            result = converter.convert(tmp_path)
            md = result.document.export_to_markdown()
            if md and md.strip():
                return md.strip()
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)
    except ImportError:
        logger.debug("Docling is not installed. Using native lightweight parsers.")
    except Exception as e:
        logger.warning(f"Docling conversion failed for {filename}: {e}. Falling back to native parsers.")
    return None


class CleanHTMLTextExtractor(HTMLParser):
    """Native standard library HTML-to-text parser."""
    def __init__(self):
        super().__init__()
        self.result = []
        self._ignore = False
        self._block_tags = {
            "p", "h1", "h2", "h3", "h4", "h5", "h6",
            "div", "article", "section", "li", "tr", "br", "blockquote"
        }

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "noscript", "svg", "head"):
            self._ignore = True
        elif tag in self._block_tags:
            self.result.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style", "noscript", "svg", "head"):
            self._ignore = False
        elif tag in self._block_tags:
            self.result.append("\n")

    def handle_data(self, data):
        if not self._ignore:
            text = data.strip()
            if text:
                self.result.append(text + " ")

    def get_text(self) -> str:
        raw = "".join(self.result)
        cleaned = re.sub(r"\n{3,}", "\n\n", raw)
        return cleaned.strip()


def extract_text_from_html(content_bytes: bytes) -> str:
    """
    Extract readable text from HTML content, stripping scripts, CSS, and boilerplate.
    Uses BeautifulSoup if available, or native html.parser.
    """
    # 1. Try BeautifulSoup if installed
    try:
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(content_bytes, "html.parser")
        for unwanted in soup(["script", "style", "noscript", "svg", "header", "footer", "nav"]):
            unwanted.decompose()

        text = soup.get_text(separator="\n", strip=True)
        if text:
            return text
    except ImportError:
        pass
    except Exception as e:
        logger.warning(f"BeautifulSoup HTML extraction failed: {e}")

    # 2. Native standard library fallback
    for encoding in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
        try:
            html_str = content_bytes.decode(encoding)
            extractor = CleanHTMLTextExtractor()
            extractor.feed(html_str)
            text = extractor.get_text()
            if text:
                return text
        except UnicodeDecodeError:
            continue
        except Exception as e:
            logger.warning(f"Native HTML parser failed: {e}")

    return ""


def extract_text_from_docx_fallback(content_bytes: bytes) -> str:
    """
    Extract text from a .docx file using Python standard library (zipfile + XML).
    Works without python-docx installed.
    """
    try:
        with zipfile.ZipFile(io.BytesIO(content_bytes)) as z:
            xml_content = z.read("word/document.xml")
            tree = ET.fromstring(xml_content)

            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

            paragraphs = []
            for p in tree.findall(".//w:p", ns):
                texts = [node.text for node in p.findall(".//w:t", ns) if node.text]
                if texts:
                    paragraphs.append("".join(texts))

            return "\n\n".join(paragraphs).strip()
    except Exception as e:
        logger.warning(f"Native docx fallback extraction failed: {e}")
        return ""


def extract_text_from_pptx(content_bytes: bytes) -> str:
    """
    Extract slide-by-slide text, tables, and speaker notes from a PowerPoint (.pptx) presentation
    (such as Client Presentations or KPI Decks).
    Uses python-pptx if installed, with a zero-dependency native zipfile + XML parser fallback.
    """
    # 1. Try python-pptx if installed
    try:
        import pptx  # type: ignore
        prs = pptx.Presentation(io.BytesIO(content_bytes))
        slides_output = []
        for idx, slide in enumerate(prs.slides, 1):
            slide_lines = [f"--- Slide {idx} ---"]
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.text_frame:
                    txt = shape.text_frame.text.strip()
                    if txt:
                        slide_lines.append(txt)
                if getattr(shape, "has_table", False) and shape.table:
                    for row in shape.table.rows:
                        cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                        if any(cells):
                            slide_lines.append("| " + " | ".join(cells) + " |")
            if getattr(slide, "has_notes_slide", False) and slide.notes_slide and slide.notes_slide.notes_text_frame:
                notes_txt = slide.notes_slide.notes_text_frame.text.strip()
                if notes_txt:
                    slide_lines.append(f"[Speaker Notes Slide {idx}]: {notes_txt}")
            if len(slide_lines) > 1:
                slides_output.append("\n".join(slide_lines))
        if slides_output:
            return "\n\n".join(slides_output).strip()
    except ImportError:
        pass
    except Exception as e:
        logger.warning(f"python-pptx parsing failed, using native XML fallback: {e}")

    # 2. Zero-dependency native zipfile + XML extraction for .pptx
    try:
        with zipfile.ZipFile(io.BytesIO(content_bytes)) as z:
            ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
            slide_files = [
                f for f in z.namelist()
                if re.match(r"^ppt/slides/slide\d+\.xml$", f)
            ]
            # Sort numerically by slide number
            slide_files.sort(key=lambda name: int(re.search(r"slide(\d+)\.xml$", name).group(1)))  # type: ignore

            slides_output = []
            for idx, slide_path in enumerate(slide_files, 1):
                xml_bytes = z.read(slide_path)
                tree = ET.fromstring(xml_bytes)
                slide_lines = [f"--- Slide {idx} ---"]

                for p in tree.findall(".//a:p", ns):
                    runs = [t.text for t in p.findall(".//a:t", ns) if t.text]
                    line = "".join(runs).strip()
                    if line:
                        slide_lines.append(line)

                # Also check speaker notes if present
                notes_path = f"ppt/notesSlides/notesSlide{idx}.xml"
                if notes_path in z.namelist():
                    try:
                        notes_tree = ET.fromstring(z.read(notes_path))
                        notes_Lines = []
                        for p in notes_tree.findall(".//a:p", ns):
                            runs = [t.text for t in p.findall(".//a:t", ns) if t.text]
                            n_line = "".join(runs).strip()
                            if n_line and not n_line.isdigit():
                                notes_Lines.append(n_line)
                        if notes_Lines:
                            slide_lines.append(f"[Speaker Notes Slide {idx}]: " + " ".join(notes_Lines))
                    except Exception:
                        pass

                if len(slide_lines) > 1:
                    slides_output.append("\n".join(slide_lines))

            return "\n\n".join(slides_output).strip()
    except Exception as e:
        logger.error(f"Native PPTX extraction failed: {e}")
        raise ValueError(f"Could not extract text from PowerPoint file: {e}")


def extract_rows_from_xlsx(content_bytes: bytes) -> dict[str, list[list[str]]]:
    """
    Extract worksheet tables from an Excel (.xlsx) spreadsheet as {sheet_name: [[cell, ...], ...]}.
    Uses openpyxl if installed, with a zero-dependency native zipfile + XML parser fallback.
    """
    # 1. Try openpyxl if installed
    try:
        import openpyxl  # type: ignore
        wb = openpyxl.load_workbook(io.BytesIO(content_bytes), data_only=True, read_only=True)
        sheets_data: dict[str, list[list[str]]] = {}
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows_list: list[list[str]] = []
            for row in ws.iter_rows(values_only=True):
                str_row = [str(cell).strip() if cell is not None else "" for cell in row]
                if any(str_row):
                    rows_list.append(str_row)
            if rows_list:
                sheets_data[sheet_name] = rows_list
        if sheets_data:
            return sheets_data
    except ImportError:
        pass
    except Exception as e:
        logger.warning(f"openpyxl extraction failed, using native XML fallback: {e}")

    # 2. Zero-dependency native zipfile + XML parser for .xlsx
    sheets_data = {}
    try:
        with zipfile.ZipFile(io.BytesIO(content_bytes)) as z:
            ns_main = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}

            # Read shared strings table if present
            shared_strings: list[str] = []
            if "xl/sharedStrings.xml" in z.namelist():
                ss_tree = ET.fromstring(z.read("xl/sharedStrings.xml"))
                for si in ss_tree.findall(".//main:si", ns_main):
                    texts = [t.text for t in si.findall(".//main:t", ns_main) if t.text]
                    shared_strings.append("".join(texts))

            # Map sheet names from workbook.xml
            sheet_names: list[str] = []
            if "xl/workbook.xml" in z.namelist():
                wb_tree = ET.fromstring(z.read("xl/workbook.xml"))
                for sh in wb_tree.findall(".//main:sheet", ns_main):
                    sheet_names.append(sh.attrib.get("name", f"Sheet{len(sheet_names)+1}"))

            sheet_files = sorted(
                [f for f in z.namelist() if re.match(r"^xl/worksheets/sheet\d+\.xml$", f)],
                key=lambda name: int(re.search(r"sheet(\d+)\.xml$", name).group(1)),  # type: ignore
            )

            def _col_letters_to_idx(cell_ref: str) -> int:
                letters = re.match(r"^([A-Z]+)", cell_ref.upper())
                if not letters:
                    return 0
                idx_val = 0
                for ch in letters.group(1):
                    idx_val = idx_val * 26 + (ord(ch) - ord("A") + 1)
                return max(0, idx_val - 1)

            for s_idx, sheet_path in enumerate(sheet_files):
                s_name = sheet_names[s_idx] if s_idx < len(sheet_names) else f"Sheet {s_idx + 1}"
                ws_tree = ET.fromstring(z.read(sheet_path))
                rows_list: list[list[str]] = []
                max_cols = 0

                parsed_sparse_rows: list[dict[int, str]] = []
                for row_elem in ws_tree.findall(".//main:sheetData/main:row", ns_main):
                    row_map: dict[int, str] = {}
                    for c_elem in row_elem.findall("main:c", ns_main):
                        cell_ref = c_elem.attrib.get("r", "")
                        col_idx = _col_letters_to_idx(cell_ref) if cell_ref else len(row_map)
                        c_type = c_elem.attrib.get("t", "")
                        v_elem = c_elem.find("main:v", ns_main)
                        is_elem = c_elem.find("main:is", ns_main)

                        val_str = ""
                        if c_type == "s" and v_elem is not None and v_elem.text is not None:
                            try:
                                ss_idx = int(v_elem.text)
                                if 0 <= ss_idx < len(shared_strings):
                                    val_str = shared_strings[ss_idx]
                            except ValueError:
                                val_str = v_elem.text
                        elif c_type == "inlineStr" and is_elem is not None:
                            texts = [t.text for t in is_elem.findall(".//main:t", ns_main) if t.text]
                            val_str = "".join(texts)
                        elif v_elem is not None and v_elem.text is not None:
                            val_str = v_elem.text

                        val_str = val_str.strip()
                        if val_str:
                            row_map[col_idx] = val_str
                            max_cols = max(max_cols, col_idx + 1)

                    if row_map:
                        parsed_sparse_rows.append(row_map)

                for r_map in parsed_sparse_rows:
                    dense_row = [r_map.get(c_i, "") for c_i in range(max_cols)]
                    if any(dense_row):
                        rows_list.append(dense_row)

                if rows_list:
                    sheets_data[s_name] = rows_list
    except Exception as e:
        logger.error(f"Native XLSX extraction failed: {e}")

    return sheets_data


def extract_text_from_xlsx(content_bytes: bytes) -> str:
    """
    Convert an Excel (.xlsx) spreadsheet (Inventory, KPI Dashboard, Financial Report)
    into clean Markdown tables and row summaries for high-accuracy RAG retrieval.
    """
    sheets_data = extract_rows_from_xlsx(content_bytes)
    if not sheets_data:
        raise ValueError("No readable worksheet data found in Excel (.xlsx) file.")

    output_blocks: list[str] = []
    for sheet_name, rows in sheets_data.items():
        if not rows:
            continue
        headers = rows[0]
        block_lines = [f"### Worksheet: {sheet_name}"]
        header_line = "| " + " | ".join(h if h else f"Col_{i+1}" for i, h in enumerate(headers)) + " |"
        sep_line = "| " + " | ".join(["---"] * len(headers)) + " |"
        block_lines.append(header_line)
        block_lines.append(sep_line)

        for r in rows[1:]:
            padded = r + [""] * max(0, len(headers) - len(r))
            row_md = "| " + " | ".join(padded[:len(headers)]) + " |"
            # Also add a key-value inline representation when headers are descriptive (helps BM25 & embeddings)
            kv_pairs = [
                f"{headers[i]}: {padded[i]}"
                for i in range(min(len(headers), len(padded)))
                if headers[i] and padded[i]
            ]
            if kv_pairs:
                block_lines.append(f"{row_md}  <!-- {' ; '.join(kv_pairs)} -->")
            else:
                block_lines.append(row_md)

        output_blocks.append("\n".join(block_lines))

    return "\n\n".join(output_blocks).strip()


def extract_text_from_pdf(content_bytes: bytes) -> str:
    """
    Extract text from a PDF file using pypdf, pypdf2, or pymupdf if available.
    """
    # 1. Try pypdf (recommended modern library)
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(content_bytes))
        pages_text = []
        for idx, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                pages_text.append(f"--- Page {idx + 1} ---\n{text.strip()}")
        return "\n\n".join(pages_text).strip()
    except ImportError:
        pass
    except Exception as e:
        logger.error(f"pypdf extraction failed: {e}")

    # 2. Try legacy pypdf2
    try:
        import PyPDF2
        reader = PyPDF2.PdfReader(io.BytesIO(content_bytes))
        pages_text = [page.extract_text() or "" for page in reader.pages]
        return "\n\n".join(t.strip() for t in pages_text if t.strip())
    except ImportError:
        pass
    except Exception as e:
        logger.error(f"PyPDF2 extraction failed: {e}")

    # 3. Try pymupdf / fitz
    try:
        import fitz
        doc = fitz.open(stream=content_bytes, filetype="pdf")
        pages_text = [page.get_text() for page in doc]
        return "\n\n".join(t.strip() for t in pages_text if t.strip())
    except ImportError:
        pass
    except Exception as e:
        logger.error(f"fitz extraction failed: {e}")

    raise RuntimeError(
        "PDF extraction requires 'pypdf'. Please run in backend: 'pip install pypdf'"
    )


def extract_text_from_docx(content_bytes: bytes) -> str:
    """
    Extract text from a .docx file using python-docx if installed,
    falling back to native zipfile + XML extraction.
    """
    # 1. Try python-docx
    try:
        import docx
        doc = docx.Document(io.BytesIO(content_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    paragraphs.append(" | ".join(row_texts))
        extracted = "\n\n".join(paragraphs).strip()
        if extracted:
            return extracted
    except ImportError:
        pass
    except Exception as e:
        logger.warning(f"python-docx parsing error, falling back to native: {e}")

    # 2. Fallback to native zipfile/XML
    fallback_text = extract_text_from_docx_fallback(content_bytes)
    if fallback_text:
        return fallback_text

    raise RuntimeError(
        "DOCX extraction requires 'python-docx'. Please run in backend: 'pip install python-docx'"
    )


def parse_document(filename: str, content_bytes: bytes) -> Tuple[str, str]:
    """
    Parse document content based on file extension.
    Prioritizes IBM Docling if installed, with zero-dependency native fallbacks:
    - .pdf: Docling -> pypdf -> fitz
    - .docx / .doc: Docling -> python-docx -> native zipfile/XML
    - .pptx: Docling -> python-pptx -> native zipfile/XML (Slide-by-Slide + Notes)
    - .xlsx: openpyxl -> native zipfile/XML (Markdown Tables + Key-Value Rows)
    - .html / .htm: Docling -> BeautifulSoup -> native html.parser
    - .txt / .md / .csv / .json / .log: text decoders
    Returns: (extracted_text, detected_type)
    """
    ext = os.path.splitext(filename)[1].lower()

    # 1. Try Docling for rich documents if installed
    if ext in [".pdf", ".docx", ".doc", ".pptx", ".html", ".htm"]:
        docling_result = try_docling_parse(filename, content_bytes)
        if docling_result:
            logger.info(f"Successfully parsed {filename} with Docling ({len(docling_result)} chars).")
            return docling_result, ext[1:]

    # 2. Format-specific parsers
    if ext == ".pdf":
        text = extract_text_from_pdf(content_bytes)
        return text, "pdf"

    elif ext in [".docx", ".doc"]:
        text = extract_text_from_docx(content_bytes)
        return text, "docx"

    elif ext == ".pptx":
        text = extract_text_from_pptx(content_bytes)
        if not text:
            raise ValueError(f"Could not extract readable slide text from PowerPoint file '{filename}'.")
        return text, "pptx"

    elif ext == ".xlsx":
        text = extract_text_from_xlsx(content_bytes)
        return text, "xlsx"

    elif ext in [".html", ".htm"]:
        text = extract_text_from_html(content_bytes)
        if not text:
            raise ValueError(f"Could not extract readable text from HTML file '{filename}'.")
        return text, "html"

    elif ext in [".txt", ".md", ".markdown", ".json", ".csv", ".log"]:
        for encoding in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
            try:
                text = content_bytes.decode(encoding)
                return text, ext[1:] if ext else "text"
            except UnicodeDecodeError:
                continue
        raise ValueError(f"Unable to decode text file '{filename}' with supported encodings.")

    else:
        # Best effort decode as plain text
        try:
            text = content_bytes.decode("utf-8")
            return text, "text"
        except UnicodeDecodeError:
            raise ValueError(
                f"Unsupported file format '{ext}'. Supported formats: .pdf, .docx, .pptx, .xlsx, .html, .htm, .txt, .md, .csv, .json"
            )
